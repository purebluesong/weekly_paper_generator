# coding=utf-8
from docx import Document

__author__ = 'Sprout'
__doc__ = u'''
    提供转化 data 为 docx 的类
'''


class FinalPaperWeekPlan:
    def __init__(self, cells=None, title=None, template='sample.docx'):
        u'''
            params:
                cells: {} 允许使用者自己定义自己的模板中的属性
                title: '' 允许使用者定义自己生成文档的题目
                template: '' 允许使用者定义自己的模板, 默认为 sample.docx
        '''
        self.report_template = Document(template)
        self.table = self.report_template.tables[0]
        self.title_template = title or u"软件学院本科实习&毕业设计（论文）周报-%s-%s.docx"
        cells = cells or {
            "name": (1, 1),
            "tel": (1, 3),
            "email": (1, 6),
            "base": (2, 1),
            "mentor": (2, 4),
            "inner_mentor": (2, 8),
            "report_time": (3, 1),
            "week": (3, 4),
            "duration": (4, 1),
            "point": (5, 1),
            "content": (6, 1),
            "problem": (7, 1),
            "summary": (8, 1),
            "nextweekplan": (9, 1),
        }
        self.cells = {}
        for k, v in cells.iteritems():
            row, col = v
            self.cells[k] = self.table.cell(row, col)

    def convert(self, data={}, destnation='', titleparams=()):
        u'''
            params:
                data :
                {
                    "name": value,
                    "tel": value,
                    "email": value,
                    "base": value,
                    "mentor": value,
                    "inner_mentor": value,
                    "report_time": value,
                    "week": value,
                    "duration": value,
                    "point": value,
                    "content": value,
                    "problem": value,
                    "summary": value,
                    "nextweekplan": value,
                }
                destnation :
                the path doc will be store, end with '/'
        '''
        if not data:
            raise RuntimeError('missing doc data')
        for k in self.cells:
            self.cells[k].text = data.get(k) or self.cells[k].text

        title = self.title_template % titleparams
        self.report_template.save(destnation+title)
